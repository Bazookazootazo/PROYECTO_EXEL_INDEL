import os
import calendar
from datetime import datetime, date, timedelta
import customtkinter as ctk
from tkinter import filedialog, messagebox
import mysql.connector

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

from utilerias.paleta_colores import colores, fuentes


# Minutos de tolerancia en la ENTRADA (Contrato Colectivo de Trabajo)
TOLERANCIA_MIN = 10


class VentanaGenerarReporte(ctk.CTkToplevel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.parent = parent

        self.title("Generación de Reportes Quincenales - Control de Acceso")
        self.configure(fg_color=colores.BG_VENTANA)

        # Configuración de pantalla completa
        self.state('zoomed')
        self.resizable(True, True)

        self.focus()
        self.lift()

        self._crear_interfaz()
        self.cargar_areas()

    # --- CONEXIÓN BASE DE DATOS ---
    @staticmethod
    def obtener_conexion_bd():
        return mysql.connector.connect(
            host="localhost",
            user="root",
            password="",
            database="control_asistencia"
        )

    def cargar_areas(self):
        """Carga la lista de áreas desde la BD al combobox."""
        try:
            conn = self.obtener_conexion_bd()
            cursor = conn.cursor()
            cursor.execute("SELECT DISTINCT area FROM profesores WHERE area IS NOT NULL AND area != '' ORDER BY area ASC")
            filas = cursor.fetchall()
            areas = [f[0] for f in filas]

            if areas:
                self.cb_area.configure(values=areas)
                self.cb_area.set(areas[0])
            else:
                self.cb_area.configure(values=["Sin áreas registradas"])
                self.cb_area.set("Sin áreas registradas")

            cursor.close()
            conn.close()
        except Exception as e:
            messagebox.showerror("Error BD", f"No se pudieron cargar las áreas:\n{e}", parent=self)

    # --- LÓGICA DE CÁLCULO Y GENERACIÓN DE EXCEL ---

    def obtener_rango_quincena(self, anio, mes, quincena_num):
        """Devuelve las fechas de inicio y fin de la quincena elegida."""
        if quincena_num == 1:
            fecha_inicio = date(anio, mes, 1)
            fecha_fin = date(anio, mes, 15)
        else:
            _, ultimo_dia = calendar.monthrange(anio, mes)
            fecha_inicio = date(anio, mes, 16)
            fecha_fin = date(anio, mes, ultimo_dia)
        return fecha_inicio, fecha_fin

    def generar_excel_quincenal(self):
        area_seleccionada = self.cb_area.get().strip()
        if not area_seleccionada or area_seleccionada == "Sin áreas registradas":
            messagebox.showwarning("Advertencia", "Selecciona un área válida.", parent=self)
            return

        try:
            anio = int(self.cb_anio.get())
            mes_idx = self.meses.index(self.cb_mes.get()) + 1
            quincena_num = 1 if "1ra" in self.cb_quincena.get() else 2
        except Exception:
            messagebox.showerror("Error", "Selección de periodo no válida.", parent=self)
            return

        fecha_inicio, fecha_fin = self.obtener_rango_quincena(anio, mes_idx, quincena_num)

        # Pedir ubicación para guardar
        nombre_defecto = f"Reporte_Asistencia_{area_seleccionada.replace(' ', '_')}_{fecha_inicio.strftime('%Y%m%d')}_{fecha_fin.strftime('%Y%m%d')}.xlsx"
        ruta_guardado = filedialog.asksaveasfilename(
            title="Guardar Reporte Excel",
            initialfile=nombre_defecto,
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            parent=self
        )

        if not ruta_guardado:
            return

        try:
            self.lbl_estado_texto.configure(text="Generando reporte...")
            self.lbl_estado_dot.configure(text_color=colores.ESTADO_PROCESANDO)
            self.update()

            # Consultar Datos de la BD
            conn = self.obtener_conexion_bd()
            cursor = conn.cursor(dictionary=True)

            query_docentes = """
                SELECT id_profesor, nombre 
                FROM profesores 
                WHERE area = %s 
                ORDER BY nombre ASC
            """
            cursor.execute(query_docentes, (area_seleccionada,))
            docentes = cursor.fetchall()

            if not docentes:
                messagebox.showwarning("Sin Datos", f"No hay docentes registrados en el área: {area_seleccionada}", parent=self)
                return

            # RESTRICCIÓN: Cambia a True si deseas quitar los fines de semana para que las columnas se vean más grandes y legibles al imprimir
            OBLIGAR_SOLO_DIAS_HABILES = True 

            dias_quincena = []
            curr = fecha_inicio
            while curr <= fecha_fin:
                if not OBLIGAR_SOLO_DIAS_HABILES or curr.weekday() < 5:  # Lunes a Viernes
                    dias_quincena.append(curr)
                curr += timedelta(days=1)

            # Construir Libro Excel
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Control Quincenal"
            ws.views.sheetView[0].showGridLines = True

            # --- CONFIGURACIÓN DE PÁGINA PARA IMPRESIÓN Y PDF PERFECTO ---
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
            ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            
            ws.page_margins.left = 0.3
            ws.page_margins.right = 0.3
            ws.page_margins.top = 0.4
            ws.page_margins.bottom = 0.4

            # --- ESTILOS INSTITUCIONALES ---
            f_titulo = Font(name="Arial", size=10, bold=True)
            f_subtitulo = Font(name="Arial", size=9, bold=True)
            f_italic = Font(name="Arial", size=7.5, italic=True)
            f_bold = Font(name="Arial", size=7.5, bold=True)
            f_normal = Font(name="Arial", size=7)
            f_nota = Font(name="Arial", size=6.5, bold=True)

            fill_header = PatternFill(start_color="FF000000", end_color="FF000000", fill_type="solid")
            f_header = Font(name="Arial", size=6.5, bold=True, color="FFFFFFFF")
            f_sub_header = Font(name="Arial", size=5.5, bold=False, color="FFFFFFFF")

            thin_border = Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )

            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            align_left = Alignment(horizontal="left", vertical="center")

            total_cols = len(dias_quincena) * 2
            col_fin_letter = get_column_letter(1 + total_cols)
            col_total_letter = get_column_letter(2 + total_cols)

            ws.row_dimensions[1].height = 42

            # --- INSERTAR LOGOTIPOS CON RUTA ABSOLUTA GARANTIZADA ---
            script_dir = os.path.dirname(os.path.abspath(__file__))

            for img_name in ["logo_utn.png", "Imagen2.png"]:
                ruta_img = os.path.join(script_dir, img_name)
                if os.path.exists(ruta_img):
                    try:
                        img_utn = OpenpyxlImage(ruta_img)
                        img_utn.width = 40
                        img_utn.height = 40
                        ws.add_image(img_utn, "A1")
                        break
                    except Exception:
                        pass

            for img_name in ["logo_iso.png", "Imagen.png"]:
                ruta_img = os.path.join(script_dir, img_name)
                if os.path.exists(ruta_img):
                    try:
                        img_iso = OpenpyxlImage(ruta_img)
                        img_iso.width = 40
                        img_iso.height = 40
                        ws.add_image(img_iso, f"{col_total_letter}1")
                        break
                    except Exception:
                        pass

            # --- ENCABEZADOS DE DOCUMENTO ---
            ws.merge_cells(f"A1:{col_total_letter}1")
            ws["A1"] = "UNIVERSIDAD TECNOLÓGICA DE NOGALES, SONORA\nDirección de Administración y Finanzas\nDepartamento de Recursos Humanos\nControl Quincenal de Asistencia Administrativo y Profesor de Tiempo Completo"
            ws["A1"].font = Font(name="Arial", size=8, bold=True)
            ws["A1"].alignment = align_center

            ws.merge_cells(f"A3:{col_total_letter}3")
            ws["A3"] = f"Área/Dirección: {area_seleccionada}"
            ws["A3"].font = f_subtitulo
            ws["A3"].alignment = align_center

            ws.merge_cells(f"A5:{col_total_letter}5")
            ws["A5"] = f"Quincena del: {fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}"
            ws["A5"].font = f_italic
            ws["A5"].alignment = align_center

            # --- TABLA DE DATOS (Encabezados Fila 7 a 9) ---
            header_start_row = 7
            header_end_row = 9

            for r in range(header_start_row, header_end_row + 1):
                for c in range(1, 3 + total_cols):
                    celda = ws.cell(row=r, column=c)
                    celda.fill = fill_header
                    celda.font = f_header
                    celda.alignment = align_center
                    celda.border = thin_border

            ws.merge_cells(f"A{header_start_row}:A{header_end_row}")
            ws[f"A{header_start_row}"] = "NOMBRE"

            ws.merge_cells(f"B{header_start_row}:{col_fin_letter}{header_start_row}")
            ws["B7"] = "DÍAS / ENTRADAS Y SALIDAS"

            ws.merge_cells(f"{col_total_letter}{header_start_row}:{col_total_letter}{header_end_row}")
            ws[f"{col_total_letter}{header_start_row}"] = "TOTAL"

            curr_col = 2
            for d in dias_quincena:
                c1_letra = get_column_letter(curr_col)
                c2_letra = get_column_letter(curr_col + 1)

                ws.merge_cells(f"{c1_letra}8:{c2_letra}8")
                ws[f"{c1_letra}8"] = d.day
                ws[f"{c1_letra}8"].font = f_header

                ws[f"{c1_letra}9"] = "Entrada"
                ws[f"{c1_letra}9"].font = f_sub_header
                ws[f"{c2_letra}9"] = "Salida"
                ws[f"{c2_letra}9"].font = f_sub_header

                curr_col += 2

            # --- PROCESAR FILAS DE DOCENTES ---
            fila_actual = 10

            for doc in docentes:
                id_prof = doc['id_profesor']
                nombre_prof = doc['nombre']

                ws.cell(row=fila_actual, column=1, value=nombre_prof).font = f_bold
                ws.cell(row=fila_actual, column=1).alignment = align_left
                ws.cell(row=fila_actual, column=1).border = thin_border

                total_retardo_min = 0
                col_dia = 2

                for d in dias_quincena:
                    st_date = datetime.combine(d, datetime.min.time())
                    end_date = datetime.combine(d, datetime.max.time())
                    id_dia_bd = d.weekday() + 1

                    q_horario = "SELECT hora_entrada, hora_salida FROM horarios_docentes WHERE id_profesor = %s AND id_dia = %s"
                    cursor.execute(q_horario, (id_prof, id_dia_bd))
                    horario_doc = cursor.fetchone()

                    q_checadas = "SELECT fecha_hora FROM checadas WHERE id_profesor = %s AND fecha_hora BETWEEN %s AND %s ORDER BY fecha_hora ASC"
                    cursor.execute(q_checadas, (id_prof, st_date, end_date))
                    checadas = [row['fecha_hora'] for row in cursor.fetchall()]

                    val_entrada = ""
                    val_salida = ""

                    if not horario_doc:
                        val_entrada = "N/A"
                        val_salida = "N/A"
                    elif not checadas:
                        val_entrada = "FI"
                        val_salida = "FI"
                    else:
                        h_entrada_esp = horario_doc['hora_entrada']
                        h_salida_esp = horario_doc['hora_salida']

                        dt_entrada_esp = datetime.combine(d, (datetime.min + h_entrada_esp).time() if isinstance(h_entrada_esp, timedelta) else h_entrada_esp)
                        dt_salida_esp = datetime.combine(d, (datetime.min + h_salida_esp).time() if isinstance(h_salida_esp, timedelta) else h_salida_esp)
                        midpoint = dt_entrada_esp + (dt_salida_esp - dt_entrada_esp) / 2

                        if len(checadas) == 1:
                            p = checadas[0]
                            primera_ch, ultima_ch = (p, None) if p <= midpoint else (None, p)
                        else:
                            primera_ch, ultima_ch = checadas[0], checadas[-1]

                        if primera_ch:
                            primera_min = primera_ch.replace(second=0, microsecond=0)
                            limite_tolerancia = dt_entrada_esp + timedelta(minutes=TOLERANCIA_MIN)
                            if primera_min <= limite_tolerancia:
                                val_entrada = primera_ch.strftime("%H:%M")
                            else:
                                retardo_min = int((primera_min - limite_tolerancia).total_seconds() // 60)
                                val_entrada = f"{primera_ch.strftime('%H:%M')} R ({retardo_min}m)"
                                total_retardo_min += retardo_min
                        else:
                            val_entrada = "FIE"

                        if ultima_ch:
                            val_salida = ultima_ch.strftime("%H:%M")
                        else:
                            val_salida = "FIS"

                    c_ent = ws.cell(row=fila_actual, column=col_dia, value=val_entrada)
                    c_sal = ws.cell(row=fila_actual, column=col_dia + 1, value=val_salida)

                    c_ent.font = f_normal
                    c_sal.font = f_normal
                    c_ent.alignment = align_center
                    c_sal.alignment = align_center
                    c_ent.border = thin_border
                    c_sal.border = thin_border

                    col_dia += 2

                c_tot = ws.cell(row=fila_actual, column=col_dia, value=f"{total_retardo_min} min" if total_retardo_min > 0 else "0")
                c_tot.font = f_bold
                c_tot.alignment = align_center
                c_tot.border = thin_border

                fila_actual += 1

            # --- PIE DE PÁGINA Y LEYENDAS ---
            fila_actual += 2
            ws.merge_cells(f"A{fila_actual}:{col_total_letter}{fila_actual}")
            ws[f"A{fila_actual}"] = "En este Control Quincenal aparecen las entradas registradas después del horario estipulado."
            ws[f"A{fila_actual}"].font = f_normal
            ws[f"A{fila_actual}"].alignment = align_center

            fila_actual += 1
            ws.merge_cells(f"A{fila_actual}:{col_total_letter}{fila_actual}")
            ws[f"A{fila_actual}"] = "Se consideran 10 minutos de TOLERANCIA solo en el registro de las Entradas, según el Contrato Colectivo de Trabajo que actualmente nos rige"
            ws[f"A{fila_actual}"].font = f_bold
            ws[f"A{fila_actual}"].alignment = align_center

            fila_actual += 1
            ws.merge_cells(f"A{fila_actual}:{col_total_letter}{fila_actual}")
            ws[f"A{fila_actual}"] = "Después de los 10 minutos de TOLERANCIA se comenzarán a contar los minutos de Retardo del Personal de UTN."
            ws[f"A{fila_actual}"].font = f_normal
            ws[f"A{fila_actual}"].alignment = align_center

            fila_actual += 2
            ws.merge_cells(f"A{fila_actual}:{col_total_letter}{fila_actual}")
            ws[f"A{fila_actual}"] = "FI = Falta Injustificada    I = Incapacidad    AM = Atención Médica    P = Permiso    C = Comisión    FIS = Salida No Registrada    FIE = Entrada No Registrada"
            ws[f"A{fila_actual}"].font = f_nota
            ws[f"A{fila_actual}"].alignment = align_center

            # --- BLOQUE INFERIOR (F01PGRRH02.01) ---
            fila_actual += 3
            ws.cell(row=fila_actual, column=1, value="F01PGRRH02.01").font = Font(name="Arial", size=6.5, bold=True)
            ws.cell(row=fila_actual, column=1).alignment = align_left

            c_rev_col1 = get_column_letter(total_cols - 1)
            c_rev_col2 = get_column_letter(total_cols)
            c_rev_col3 = get_column_letter(total_cols + 1)
            c_rev_col4 = get_column_letter(total_cols + 2)

            ws.merge_cells(f"{c_rev_col1}{fila_actual}:{c_rev_col1}{fila_actual+1}")
            ws[f"{c_rev_col1}{fila_actual}"] = "Fecha revisión\n02/04/2025"
            ws[f"{c_rev_col1}{fila_actual}"].font = Font(name="Arial", size=5, bold=True)
            ws[f"{c_rev_col1}{fila_actual}"].alignment = align_center

            ws.merge_cells(f"{c_rev_col2}{fila_actual}:{c_rev_col3}{fila_actual+1}")
            ws[f"{c_rev_col2}{fila_actual}"] = "Hoja 1 de 1"
            ws[f"{c_rev_col2}{fila_actual}"].font = Font(name="Arial", size=5.5, bold=True)
            ws[f"{c_rev_col2}{fila_actual}"].alignment = align_center

            ws.merge_cells(f"{c_rev_col4}{fila_actual}:{c_rev_col4}{fila_actual+1}")
            ws[f"{c_rev_col4}{fila_actual}"] = "Rev. 01"
            ws[f"{c_rev_col4}{fila_actual}"].font = Font(name="Arial", size=5.5, bold=True)
            ws[f"{c_rev_col4}{fila_actual}"].alignment = align_center

            for r_idx in range(fila_actual, fila_actual + 2):
                for col_let in [c_rev_col1, c_rev_col2, c_rev_col3, c_rev_col4]:
                    ws[f"{col_let}{r_idx}"].border = thin_border
                    ws[f"{col_let}{r_idx}"].fill = fill_header
                    ws[f"{col_let}{r_idx}"].font = Font(name="Arial", size=5.5, bold=True, color="FFFFFFFF")

            ws.column_dimensions['A'].width = 22
            for col_i in range(2, 2 + total_cols + 2):
                col_letter = get_column_letter(col_i)
                ws.column_dimensions[col_letter].width = 6.0

            wb.save(ruta_guardado)
            cursor.close()
            conn.close()

            self.lbl_estado_texto.configure(text="Reporte generado exitosamente.")
            self.lbl_estado_dot.configure(text_color=colores.ESTADO_EXITO)
            messagebox.showinfo("Éxito", f"El reporte ha sido generado correctamente en:\n{ruta_guardado}", parent=self)

        except Exception as err:
            self.lbl_estado_texto.configure(text="Error al generar reporte")
            self.lbl_estado_dot.configure(text_color=colores.ESTADO_ERROR)
            messagebox.showerror("Error", f"Fallo durante la generación del reporte:\n{err}", parent=self)

    def regresar_al_menu(self):
        if self.parent:
            self.parent.wm_state('normal')
            self.parent.deiconify()
        self.destroy()

    # --- DISEÑO DE LA INTERFAZ ---
    def _crear_interfaz(self):
        main_card = ctk.CTkFrame(
            master=self,
            corner_radius=20,
            fg_color=colores.BG_PANEL,
            border_width=1,
            border_color=colores.BORDE
        )
        main_card.pack(fill="both", expand=True, padx=15, pady=15)

        franja_acento = ctk.CTkFrame(master=main_card, corner_radius=0, fg_color=colores.ACENTO, height=4)
        franja_acento.pack(fill="x", side="top")

        header_frame = ctk.CTkFrame(master=main_card, fg_color=colores.BG_PANEL_ALT, corner_radius=15, height=90, border_width=1, border_color=colores.BORDE_SUAVE)
        header_frame.pack(fill="x", padx=25, pady=(20, 15))

        icono_bg = ctk.CTkFrame(master=header_frame, fg_color="#352A12", corner_radius=14, width=58, height=58)
        icono_bg.pack(side="left", padx=(22, 15), pady=15)
        icono_bg.pack_propagate(False)

        lbl_icon = ctk.CTkLabel(master=icono_bg, text="📋", font=ctk.CTkFont(size=26))
        lbl_icon.pack(expand=True)

        title_box = ctk.CTkFrame(master=header_frame, fg_color="transparent")
        title_box.pack(side="left", fill="y", pady=15)

        lbl_titulo = ctk.CTkLabel(
            master=title_box,
            text="Generador de Reporte Quincenal de Asistencia",
            font=fuentes.titulo(size=22),
            text_color=colores.TEXTO_PRIMARIO,
        )
        lbl_titulo.pack(anchor="w")

        lbl_subtitulo = ctk.CTkLabel(
            master=title_box,
            text="Selecciona el área/dirección y el periodo quincenal para exportar el archivo de control en Excel.",
            font=fuentes.cuerpo(size=13),
            text_color=colores.TEXTO_SECUNDARIO,
        )
        lbl_subtitulo.pack(anchor="w", pady=(2, 0))

        params_frame = ctk.CTkFrame(master=main_card, fg_color=colores.BG_PANEL_ALT, corner_radius=15, border_width=1, border_color=colores.BORDE_SUAVE)
        params_frame.pack(fill="x", padx=25, pady=10, ipady=15)

        lbl_sec_params = ctk.CTkLabel(master=params_frame, text="PARÁMETROS DEL REPORTE", font=fuentes.pequena(size=11), text_color=colores.TEXTO_SECUNDARIO)
        lbl_sec_params.pack(anchor="w", padx=25, pady=(15, 10))

        box_inputs = ctk.CTkFrame(master=params_frame, fg_color="transparent")
        box_inputs.pack(fill="x", padx=25)
        box_inputs.grid_columnconfigure((0, 1, 2, 3), weight=1, uniform="col")

        lbl_area = ctk.CTkLabel(master=box_inputs, text="Área / Dirección:", font=fuentes.seccion(size=13), text_color=colores.TEXTO_PRIMARIO)
        lbl_area.grid(row=0, column=0, sticky="w", pady=(0, 5))
        self.cb_area = ctk.CTkComboBox(master=box_inputs, font=fuentes.cuerpo(size=13), height=38, corner_radius=8)
        self.cb_area.grid(row=1, column=0, sticky="ew", padx=(0, 10))

        lbl_anio = ctk.CTkLabel(master=box_inputs, text="Año:", font=fuentes.seccion(size=13), text_color=colores.TEXTO_PRIMARIO)
        lbl_anio.grid(row=0, column=1, sticky="w", pady=(0, 5))
        anio_actual = datetime.now().year
        self.cb_anio = ctk.CTkComboBox(master=box_inputs, values=[str(y) for y in range(anio_actual - 2, anio_actual + 2)], font=fuentes.cuerpo(size=13), height=38, corner_radius=8)
        self.cb_anio.set(str(anio_actual))
        self.cb_anio.grid(row=1, column=1, sticky="ew", padx=5)

        lbl_mes = ctk.CTkLabel(master=box_inputs, text="Mes:", font=fuentes.seccion(size=13), text_color=colores.TEXTO_PRIMARIO)
        lbl_mes.grid(row=0, column=2, sticky="w", pady=(0, 5))
        self.meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        self.cb_mes = ctk.CTkComboBox(master=box_inputs, values=self.meses, font=fuentes.cuerpo(size=13), height=38, corner_radius=8)
        self.cb_mes.set(self.meses[datetime.now().month - 1])
        self.cb_mes.grid(row=1, column=2, sticky="ew", padx=5)

        lbl_quincena = ctk.CTkLabel(master=box_inputs, text="Quincena:", font=fuentes.seccion(size=13), text_color=colores.TEXTO_PRIMARIO)
        lbl_quincena.grid(row=0, column=3, sticky="w", pady=(0, 5))
        self.cb_quincena = ctk.CTkComboBox(master=box_inputs, values=["1ra Quincena (Día 01 al 15)", "2da Quincena (Día 16 al Fin)"], font=fuentes.cuerpo(size=13), height=38, corner_radius=8)
        self.cb_quincena.set("1ra Quincena (Día 01 al 15)")
        self.cb_quincena.grid(row=1, column=3, sticky="ew", padx=(10, 0))

        btn_frame = ctk.CTkFrame(master=main_card, fg_color="transparent")
        btn_frame.pack(fill="x", padx=25, pady=20)

        btn_generar = ctk.CTkButton(
            master=btn_frame,
            text="📊 Exportar Reporte a Excel",
            fg_color="#EEA942",
            hover_color="#AA7526",
            font=fuentes.boton(size=14),
            corner_radius=10,
            height=46,
            command=self.generar_excel_quincenal
        )
        btn_generar.pack(side="left", padx=(0, 15), expand=True, fill="x")

        btn_volver = ctk.CTkButton(
            master=btn_frame,
            text="Volver al Menú Principal",
            fg_color=colores.NEUTRO,
            hover_color=colores.NEUTRO_HOVER,
            font=fuentes.boton(size=14),
            corner_radius=10,
            width=220,
            height=46,
            command=self.regresar_al_menu
        )
        btn_volver.pack(side="right")

        status_bar = ctk.CTkFrame(master=main_card, fg_color=colores.BG_ENTRADA, corner_radius=10, height=40, border_width=1, border_color=colores.BORDE_SUAVE)
        status_bar.pack(fill="x", padx=25, side="bottom", pady=(0, 20))

        self.lbl_estado_dot = ctk.CTkLabel(master=status_bar, text="●", font=ctk.CTkFont(size=16), text_color=colores.ESTADO_INACTIVO)
        self.lbl_estado_dot.pack(side="left", padx=(16, 6))

        self.lbl_estado_texto = ctk.CTkLabel(
            master=status_bar,
            text="Listo para generar reporte.",
            font=fuentes.cuerpo(size=12),
            text_color=colores.TEXTO_SECUNDARIO,
        )
        self.lbl_estado_texto.pack(side="left")


if __name__ == "__main__":
    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")
    root = ctk.CTk()
    root.withdraw()
    ventana = VentanaGenerarReporte(root)
    ventana.protocol("WM_DELETE_WINDOW", root.destroy)
    root.mainloop()